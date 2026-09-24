import datetime
import pytest
from app import create_app
from models import db, User, Department, DailyReport

DATE = datetime.date(2026, 4, 1)
FROM_DATE = "2026-01-01"
TO_DATE = "2026-12-31"


class TestConfig:
    TESTING = True
    WTF_CSRF_ENABLED = False
    SQLALCHEMY_DATABASE_URI = 'sqlite:///:memory:'
    SECRET_KEY = 'test_secret_key'
    SQLALCHEMY_TRACK_MODIFICATIONS = False

@pytest.fixture
def app():
    app = create_app(TestConfig)
    with app.app_context():
        db.create_all()
        yield app
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    with app.app_context():
        yield app.test_client()


def ensure_user(username, role='admin', password='pass'):
    u = User.query.filter_by(username=username).first()
    if not u:
        u = User(username=username, role=role)
        u.set_password(password)
        db.session.add(u)
        db.session.commit()
    return u


def ensure_department(name='Терапевтичне', bed_profile_name='Терапія', bed_capacity=30, row_no=1):
    d = Department.query.filter_by(name=name).first()
    if not d:
        d = Department(
            name=name,
            bed_profile_name=bed_profile_name,
            bed_capacity=bed_capacity,
            row_no=row_no
        )
        db.session.add(d)
        db.session.commit()
    return d


def login(client, username='admin', password='pass'):
    client.post('/login', data={'username': username, 'password': password}, follow_redirects=True)


def test_form016_hospital_wide_view(app, client):
    """Test standard Form 016 hospital-wide view aggregates daily reports by month."""
    with app.app_context():
        u = ensure_user('admin')
        dept1 = ensure_department('Терапевтичне', 'Терапія', bed_capacity=30, row_no=1)
        dept2 = ensure_department('Хірургічне', 'Хірургія', bed_capacity=20, row_no=2)

        # Create some daily reports on the first day of April
        dr1 = DailyReport(
            report_date=datetime.date(2026, 4, 1),
            department_id=dept1.id,
            beds_total=30,
            patients_start=15,
            admitted_total=5,
            admitted_rural=2,
            admitted_children=1,
            discharged_total=3,
            patients_end=17,
            created_by=u.id
        )
        dr2 = DailyReport(
            report_date=datetime.date(2026, 4, 1),
            department_id=dept2.id,
            beds_total=20,
            patients_start=10,
            admitted_total=2,
            admitted_rural=0,
            admitted_children=0,
            discharged_total=1,
            patients_end=11,
            created_by=u.id
        )
        db.session.add_all([dr1, dr2])
        db.session.commit()

        login(client)
        response = client.get(
            '/statisty/form016',
            query_string={'from_date': FROM_DATE, 'to_date': TO_DATE}
        )
        assert response.status_code == 200
        html = response.get_data(as_text=True)

        # Hospital-wide view check
        assert 'ЗВЕДЕНА ВІДОМІСТЬ обліку руху хворих' in html
        assert 'Найменування місяців' in html
        assert 'Січень' in html
        assert 'Квітень' in html
        assert 'За півріччя' in html
        assert 'За рік' in html

        # Verify the 17-column layout (Column A + 16 numbered columns)
        assert '<td>А</td><td>1</td><td>2</td><td>3</td><td>4</td>' in html
        assert '<td>5</td><td>6</td><td>7</td><td>8</td><td>9</td>' in html
        assert '<td>10</td><td>11</td><td>12</td><td>13</td><td>14</td>' in html
        assert '<td>15</td><td>16</td>' in html
        assert '<td>17</td>' not in html
        assert '<td>18</td>' not in html
        assert '№ рядка' not in html
        assert 'Число місяця' not in html
        assert 'в т.ч. сільських' not in html


def test_form016_department_specific_view(app, client):
    """Test department-specific Form 016 view displays monthly consolidated rows."""
    with app.app_context():
        u = ensure_user('admin')
        dept = ensure_department('Терапевтичне', 'Терапія', bed_capacity=30, row_no=1)

        # Create daily reports for April 1st and April 2nd
        dr1 = DailyReport(
            report_date=datetime.date(2026, 4, 1),
            department_id=dept.id,
            beds_total=30,
            patients_start=10,
            admitted_total=3,
            discharged_total=2,
            patients_end=11,
            created_by=u.id
        )
        dr2 = DailyReport(
            report_date=datetime.date(2026, 4, 2),
            department_id=dept.id,
            beds_total=30,
            patients_start=11,
            admitted_total=1,
            discharged_total=4,
            patients_end=8,
            created_by=u.id
        )
        db.session.add_all([dr1, dr2])
        db.session.commit()

        login(client)
        response = client.get(
            '/statisty/form016',
            query_string={
                'from_date': FROM_DATE,
                'to_date': TO_DATE,
                'department_id': dept.id
            }
        )
        assert response.status_code == 200
        html = response.get_data(as_text=True)

        # Department specific view check
        assert 'ЗВЕДЕНА ВІДОМІСТЬ обліку руху хворих' in html
        assert 'по відділенню: <strong>Терапевтичне</strong>' in html
        assert 'Найменування місяців' in html
        assert 'Січень' in html
        assert 'Квітень' in html
        assert 'За півріччя' in html
        assert 'За рік' in html

        # Verify 17-column header layout
        assert '<td>А</td><td>1</td><td>2</td><td>3</td><td>4</td>' in html
        assert '<td>5</td><td>6</td><td>7</td><td>8</td><td>9</td>' in html
        assert '<td>10</td><td>11</td><td>12</td><td>13</td><td>14</td>' in html
        assert '<td>15</td><td>16</td>' in html


def test_form016_excel_export(app, client):
    """Test Excel export route return code, column count and monthly structure."""
    import io
    from openpyxl import load_workbook

    with app.app_context():
        u = ensure_user('admin')
        dept = ensure_department('Терапевтичне', 'Терапія', bed_capacity=30, row_no=1)

        # Create a daily report on April 1st to populate data
        dr = DailyReport(
            report_date=datetime.date(2026, 4, 1),
            department_id=dept.id,
            beds_total=30,
            patients_start=15,
            admitted_total=5,
            admitted_rural=2,
            admitted_children=1,
            discharged_total=3,
            patients_end=17,
            created_by=u.id
        )
        db.session.add(dr)
        db.session.commit()

        login(client)
        response = client.get(
            '/statisty/form016/export',
            query_string={
                'from_date': FROM_DATE,
                'to_date': TO_DATE,
                'department_id': dept.id
            }
        )
        assert response.status_code == 200
        assert response.headers['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'attachment; filename=' in response.headers['Content-Disposition']

        # Parse Excel file contents
        wb = load_workbook(io.BytesIO(response.data))
        ws = wb.active

        # Check total columns (should be 17, from A to Q)
        assert ws.max_column == 17

        # Check the row numbering in Row 8 (should be A, 1 to 16)
        row8_vals = [ws.cell(row=8, column=col).value for col in range(1, 18)]
        assert row8_vals == ['А', 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]

        # Column A header in Row 4 should be "Найменування місяців"
        assert ws.cell(row=4, column=1).value == "Найменування місяців"

        # Check merged cell ranges for the fixed layout
        merged_ranges = [r.coord for r in ws.merged_cells.ranges]
        assert 'E4:K4' in merged_ranges
        assert 'L4:L7' in merged_ranges
        assert ws.cell(row=4, column=12).value == "Померло\n(гр. 11)"

        # Check row values. Row 9 is Січень, Row 10 is Лютий, Row 11 is Березень, Row 12 is "За I квартал", Row 13 is Квітень (Month 4)
        row13_vals = [ws.cell(row=13, column=col).value for col in range(1, 18)]
        assert row13_vals[0] == 'Квітень'
        assert row13_vals[1] == 30  # beds_total on last day of April
        assert row13_vals[3] == 15  # patients_start on April 1st
        assert row13_vals[4] == 5   # admitted_total
        assert row13_vals[5] == 2   # admitted_rural
        assert row13_vals[6] == 1   # admitted_children
        assert row13_vals[9] == 3   # discharged_total
        assert row13_vals[12] == 0  # patients_end on last day of April (no report on 30.04)
        assert row13_vals[13] == 17 # bed_days_total (since patients_end was 17 on April 1st and 0 elsewhere)


def test_form016_annual_view_default(app, client):
    """Test that Form 016 defaults to the whole year when from_date is omitted, and beds_average is an integer."""
    with app.app_context():
        u = ensure_user('admin')
        dept = ensure_department('Терапевтичне', 'Терапія', bed_capacity=30, row_no=1)

        # Create a report in April with a decimal bed average if divided, but since we use beds_total=30, beds_average should be 30 (integer)
        dr = DailyReport(
            report_date=datetime.date(2026, 4, 1),
            department_id=dept.id,
            beds_total=30,
            patients_start=15,
            admitted_total=5,
            patients_end=17,
            created_by=u.id
        )
        db.session.add(dr)
        db.session.commit()

        login(client)
        # Omit 'from_date' to trigger the default annual view
        response = client.get('/statisty/form016')
        assert response.status_code == 200
        html = response.get_data(as_text=True)

        # It should show January and April because it shows the entire year up to December
        assert 'Січень' in html
        assert 'Квітень' in html
        assert 'За рік' in html
        assert 'Всього за період' not in html  # Since it covers December, it shows "За рік"

        # Check that beds_average is rendered as an integer (e.g. "30" instead of "30.0")
        assert '30.0' not in html
        assert 'class="text-center excel-computed">30</td>' in html



def _add_report(dept, user, day, **fields):
    db.session.add(DailyReport(report_date=day, department_id=dept.id,
                               created_by=user.id, **fields))


def test_form016_month_range(app, client):
    """A March–May range shows only those months and totals only their data."""
    with app.app_context():
        u = ensure_user('admin')
        dept = ensure_department()
        _add_report(dept, u, datetime.date(2026, 2, 10), beds_total=30, admitted_total=100)
        _add_report(dept, u, datetime.date(2026, 3, 5), beds_total=30, admitted_total=4)
        _add_report(dept, u, datetime.date(2026, 5, 20), beds_total=30, admitted_total=6)
        _add_report(dept, u, datetime.date(2026, 6, 1), beds_total=30, admitted_total=200)
        db.session.commit()

        login(client)
        response = client.get('/statisty/form016',
                              query_string={'from_date': '2026-03', 'to_date': '2026-05'})
        assert response.status_code == 200
        html = response.get_data(as_text=True)

        assert 'Березень' in html and 'Квітень' in html and 'Травень' in html
        assert 'Лютий' not in html and 'Червень' not in html
        assert 'За I квартал' not in html and 'За II квартал' not in html
        assert 'Всього за період (березень–травень)' in html
        assert 'Березень – Травень 2026' in html  # period label
        assert '(2 записів за' in html            # only reports within the range
        assert '<td class="text-center fw-bold">10</td>' in html  # 4 + 6 admitted


def test_form016_quarter_subtotal_needs_whole_quarter(app, client):
    with app.app_context():
        ensure_user('admin')
        ensure_department()
        login(client)
        html = client.get('/statisty/form016',
                          query_string={'from_date': '2026-04', 'to_date': '2026-06'}).get_data(as_text=True)
        assert 'За II квартал' in html
        assert 'За півріччя' not in html
        assert 'Всього за період (квітень–червень)' in html


def test_form016_to_month_means_end_of_month(app, client):
    with app.app_context():
        u = ensure_user('admin')
        dept = ensure_department()
        _add_report(dept, u, datetime.date(2026, 5, 31), beds_total=30, admitted_total=7)
        db.session.commit()
        login(client)
        html = client.get('/statisty/form016',
                          query_string={'from_date': '2026-05', 'to_date': '2026-05'}).get_data(as_text=True)
        assert 'Всього за період (травень)' in html
        assert '(1 записів за Травень 2026)' in html


def test_form016_range_across_years_is_clamped(app, client):
    with app.app_context():
        ensure_user('admin')
        ensure_department()
        login(client)
        html = client.get('/statisty/form016',
                          query_string={'from_date': '2026-11', 'to_date': '2027-02'}).get_data(as_text=True)
        assert 'в межах одного року' in html
        assert 'Грудень' in html and 'Січень' not in html
        assert 'Всього за період (листопад–грудень)' in html
