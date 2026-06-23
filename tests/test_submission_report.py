import os
import datetime
from io import BytesIO

os.environ.setdefault('SECRET_KEY', 'test-key')

import pytest
from openpyxl import load_workbook

from app import create_app
from models import db, User, Record

DATE = datetime.date(2026, 6, 10)
FROM_DATE = datetime.date(2026, 6, 1).isoformat()
TO_DATE = datetime.date(2026, 6, 30).isoformat()
PHYS = 'Москалик І.Т'


@pytest.fixture
def app():
    app = create_app()
    app.config['TESTING'] = True
    app.config['WTF_CSRF_ENABLED'] = False
    app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///:memory:'
    with app.app_context():
        db.create_all()
        yield app
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    with app.app_context():
        yield app.test_client()


def ensure_user(username, role='admin', password='pass'):
    if not User.query.filter_by(username=username).first():
        u = User(username=username, role=role)
        u.set_password(password)
        db.session.add(u)
        db.session.commit()
    return User.query.filter_by(username=username).first()


def make_record(user_id, physician=PHYS, submitted=False, dept='Терапія',
                history='12345', full_name='Тест Пацієнт', date=DATE):
    r = Record(
        date_of_discharge=date,
        full_name=full_name,
        discharge_department=dept,
        treating_physician=physician,
        history=history,
        k_days=1,
        discharge_status='Виписаний',
        history_submitted=submitted,
        created_by=user_id,
        created_at=datetime.datetime.utcnow(),
    )
    db.session.add(r)
    db.session.commit()
    return r


def login(client, username='admin', password='pass'):
    client.post('/login', data={'username': username, 'password': password},
                follow_redirects=True)


def drilldown(client, physician=PHYS, from_date=FROM_DATE, to_date=TO_DATE):
    return client.get('/admin/reports/submission-page/physician',
                      query_string={'physician': physician,
                                    'from_date': from_date, 'to_date': to_date})


# --------------------------------------------------------------------------- #
# HTML drill-down page
# --------------------------------------------------------------------------- #

def test_drilldown_lists_submitted_and_not_submitted(app, client):
    """Сторінка лікаря показує і здані, і незданні історії за період."""
    with app.app_context():
        u = ensure_user('admin')
        make_record(u.id, submitted=False, history='NOT-1', full_name='Незданий А')
        make_record(u.id, submitted=False, history='NOT-2', full_name='Незданий Б')
        make_record(u.id, submitted=True, history='SUB-1', full_name='Зданий В')

        login(client)
        rv = drilldown(client)
        assert rv.status_code == 200
        txt = rv.get_data(as_text=True)
        assert PHYS in txt
        for marker in ('NOT-1', 'NOT-2', 'SUB-1', 'Незданий А', 'Зданий В'):
            assert marker in txt
        # лічильники: 2 не здано / 1 здано
        assert 'Не здано' in txt and 'Здано' in txt


def test_drilldown_excludes_gynecology_and_reanimation(app, client):
    """Деталізація застосовує той самий фільтр виключень, що й зведення
    (_SUBMISSION_EXCL_DEPTS) — записи з виключених відділень не показуються."""
    with app.app_context():
        u = ensure_user('admin')
        make_record(u.id, submitted=False, dept='Терапія', history='KEEP-1')
        make_record(u.id, submitted=False, dept='гінекологія', history='EXCL-GYN')
        make_record(u.id, submitted=True, dept='реанімація', history='EXCL-REA')

        login(client)
        rv = drilldown(client)
        assert rv.status_code == 200
        txt = rv.get_data(as_text=True)
        assert 'KEEP-1' in txt
        assert 'EXCL-GYN' not in txt
        assert 'EXCL-REA' not in txt


def test_drilldown_only_within_period(app, client):
    """Записи поза межами періоду не відображаються."""
    with app.app_context():
        u = ensure_user('admin')
        make_record(u.id, history='IN-RANGE', date=DATE)
        make_record(u.id, history='OUT-RANGE', date=datetime.date(2026, 5, 10))

        login(client)
        rv = drilldown(client)
        txt = rv.get_data(as_text=True)
        assert 'IN-RANGE' in txt
        assert 'OUT-RANGE' not in txt


def test_drilldown_missing_physician_redirects(app, client):
    """Без параметра physician — редірект на зведену сторінку."""
    with app.app_context():
        ensure_user('admin')
        login(client)
        rv = client.get('/admin/reports/submission-page/physician',
                        query_string={'from_date': FROM_DATE, 'to_date': TO_DATE})
        assert rv.status_code == 302
        assert '/admin/reports/submission-page' in rv.headers['Location']


def test_operator_can_access_drilldown(app, client):
    """Роль operator має доступ до деталізації (як і до решти звітів)."""
    with app.app_context():
        u = ensure_user('op', role='operator')
        make_record(u.id, history='OP-OK')
        login(client, 'op')
        rv = drilldown(client)
        assert rv.status_code == 200
        assert 'OP-OK' in rv.get_data(as_text=True)


# --------------------------------------------------------------------------- #
# Summary page: clickable physician + Excel button
# --------------------------------------------------------------------------- #

def test_summary_page_has_drilldown_link_and_excel_button(app, client):
    """На зведеній сторінці лікар клікабельний і є кнопка Excel (всі лікарі)."""
    with app.app_context():
        u = ensure_user('admin')
        make_record(u.id, submitted=True)
        login(client)
        rv = client.get('/admin/reports/submission-page',
                        query_string={'from_date': FROM_DATE, 'to_date': TO_DATE})
        assert rv.status_code == 200
        txt = rv.get_data(as_text=True)
        assert '/admin/reports/submission-page/physician' in txt
        assert '/admin/reports/submission.xlsx' in txt


# --------------------------------------------------------------------------- #
# Excel exports
# --------------------------------------------------------------------------- #

def test_physician_xlsx_export(app, client):
    """Excel по одному лікарю: правильний mimetype і стовпець «Здано» Так/Ні."""
    with app.app_context():
        u = ensure_user('admin')
        make_record(u.id, submitted=False, history='NOT-1')
        make_record(u.id, submitted=True, history='SUB-1')

        login(client)
        rv = client.get('/admin/reports/submission/physician.xlsx',
                        query_string={'physician': PHYS,
                                      'from_date': FROM_DATE, 'to_date': TO_DATE})
        assert rv.status_code == 200
        assert 'spreadsheetml' in rv.headers['Content-Type']

        wb = load_workbook(BytesIO(rv.data))
        ws = wb.active
        header = [c.value for c in ws[1]]
        assert 'Здано' in header
        zdano_col = header.index('Здано') + 1
        values = {ws.cell(row=r, column=header.index('№ історії хвороби') + 1).value:
                  ws.cell(row=r, column=zdano_col).value
                  for r in range(2, ws.max_row + 1)}
        assert values.get('NOT-1') == 'Ні'
        assert values.get('SUB-1') == 'Так'


def test_all_doctors_xlsx_two_sheets(app, client):
    """Excel «по всіх»: листи «Зведення» + «Всі записи», цифри зведення вірні."""
    with app.app_context():
        u = ensure_user('admin')
        # Лікар А: 1 здано, 2 не здано
        make_record(u.id, physician='Лікар А', submitted=True, history='A-1')
        make_record(u.id, physician='Лікар А', submitted=False, history='A-2')
        make_record(u.id, physician='Лікар А', submitted=False, history='A-3')
        # виключене відділення не повинно потрапити (канонічний токен зі списку)
        make_record(u.id, physician='Лікар А', submitted=False,
                    dept='реанімація', history='A-EXCL')

        login(client)
        rv = client.get('/admin/reports/submission.xlsx',
                        query_string={'from_date': FROM_DATE, 'to_date': TO_DATE})
        assert rv.status_code == 200
        assert 'spreadsheetml' in rv.headers['Content-Type']

        wb = load_workbook(BytesIO(rv.data))
        assert wb.sheetnames == ['Зведення', 'Всі записи']

        summary = wb['Зведення']
        s_header = [c.value for c in summary[1]]
        assert s_header == ['Лікар', 'Здано', 'Не здано', 'Всього', '% здачі']
        # рядок лікаря А: 1 / 2 / 3
        row_a = [summary.cell(row=2, column=c).value for c in range(1, 6)]
        assert row_a[0] == 'Лікар А'
        assert row_a[1] == 1 and row_a[2] == 2 and row_a[3] == 3

        detail = wb['Всі записи']
        histories = [detail.cell(row=r, column=6).value  # № історії хвороби (6-та колонка)
                     for r in range(2, detail.max_row + 1)]
        assert 'A-1' in histories and 'A-2' in histories and 'A-3' in histories
        assert 'A-EXCL' not in histories


# --------------------------------------------------------------------------- #
# PDF (skip where WeasyPrint native deps are unavailable, e.g. Windows host)
# --------------------------------------------------------------------------- #

def test_physician_pdf_export(app, client):
    pytest.importorskip('weasyprint')
    with app.app_context():
        u = ensure_user('admin')
        make_record(u.id, submitted=False, history='NOT-1')
        make_record(u.id, submitted=True, history='SUB-1')

        login(client)
        rv = client.get('/admin/reports/submission/physician.pdf',
                        query_string={'physician': PHYS,
                                      'from_date': FROM_DATE, 'to_date': TO_DATE})
        assert rv.status_code == 200
        assert rv.headers['Content-Type'] == 'application/pdf'
        assert rv.data[:4] == b'%PDF'
