"""Tests for Record.patient_ehealth_id — ID пацієнта (ЕСОЗ), editable by admin/editor."""
import io
import datetime
import pytest
from openpyxl import load_workbook
from app import create_app
from models import db, User, Record, Department
from constants import STATUS_PROCESSING
from utils import normalize_ehealth_id

DATE = datetime.date(2026, 3, 1)
EID = 'd9a421c8-8efc-11f1-b457-765e37fe0f90'
EID2 = '0b6f2e4a-1c3d-4e5f-8a9b-0c1d2e3f4a5b'


@pytest.fixture
def app():
    app = create_app()
    app.config['TESTING'] = True
    app.config['WTF_CSRF_ENABLED'] = False
    with app.app_context():
        db.create_all()
        yield app
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    with app.app_context():
        yield app.test_client()


def ensure_user(username, role):
    u = User.query.filter_by(username=username).first()
    if not u:
        u = User(username=username, role=role)
        u.set_password('pass')
        db.session.add(u)
        db.session.commit()
    return u


def login(client, username):
    client.post('/login', data={'username': username, 'password': 'pass'}, follow_redirects=True)


def make_record(user_id, full_name='Test Patient', ehealth_id=None):
    if not Department.query.filter_by(name='DeptA').first():
        db.session.add(Department(name='DeptA'))
    r = Record(date_of_discharge=DATE, full_name=full_name, discharge_department='DeptA',
               treating_physician='Dr. Test', history='12345', k_days=5,
               discharge_status=STATUS_PROCESSING, patient_ehealth_id=ehealth_id,
               created_by=user_id)
    db.session.add(r)
    db.session.commit()
    return r


def edit_form(r, ehealth_id):
    return {
        'date_of_discharge': DATE.isoformat(),
        'full_name': r.full_name,
        'discharge_department': 'DeptA',
        'treating_physician': 'Dr. Test',
        'history': '12345',
        'k_days': '5',
        'discharge_status': STATUS_PROCESSING,
        'patient_ehealth_id': ehealth_id,
    }


# --- normalization ---

def test_normalize_ehealth_id():
    assert normalize_ehealth_id(EID) == EID
    assert normalize_ehealth_id(' ' + EID.upper() + ' ') == EID
    assert normalize_ehealth_id(EID.replace('-', '')) is None   # без дефісів — не приймаємо
    assert normalize_ehealth_id(EID[:-1] + 'z') is None
    assert normalize_ehealth_id('12345') is None


# --- editing ---

@pytest.mark.parametrize('role', ['editor', 'admin'])
def test_edit_saves_ehealth_id(app, client, role):
    with app.app_context():
        u = ensure_user(role, role)
        r = make_record(u.id)
        login(client, role)
        resp = client.post(f'/api/records/{r.id}/edit', data=edit_form(r, EID.upper()))
        assert resp.status_code == 200, resp.get_json()
        db.session.refresh(r)
        assert r.patient_ehealth_id == EID  # збережено в канонічному нижньому регістрі


def test_edit_page_saves_and_clears_ehealth_id(app, client):
    with app.app_context():
        u = ensure_user('editor', 'editor')
        r = make_record(u.id)
        login(client, 'editor')
        client.post(f'/records/{r.id}/edit', data=edit_form(r, EID))
        db.session.refresh(r)
        assert r.patient_ehealth_id == EID
        client.post(f'/records/{r.id}/edit', data=edit_form(r, ''))
        db.session.refresh(r)
        assert r.patient_ehealth_id is None


def test_edit_rejects_malformed_ehealth_id(app, client):
    with app.app_context():
        u = ensure_user('editor', 'editor')
        r = make_record(u.id)
        login(client, 'editor')
        resp = client.post(f'/api/records/{r.id}/edit', data=edit_form(r, 'not-a-uuid'))
        assert resp.status_code == 400
        assert 'формат' in resp.get_json()['error']
        db.session.refresh(r)
        assert r.patient_ehealth_id is None


def test_edit_rejects_duplicate_ehealth_id(app, client):
    with app.app_context():
        u = ensure_user('editor', 'editor')
        first = make_record(u.id, 'Перший Пацієнт', ehealth_id=EID)
        second = make_record(u.id, 'Другий Пацієнт')
        login(client, 'editor')
        resp = client.post(f'/api/records/{second.id}/edit', data=edit_form(second, EID))
        assert resp.status_code == 400
        assert f'#{first.id}' in resp.get_json()['error']
        db.session.refresh(second)
        assert second.patient_ehealth_id is None

        # Повторне збереження того самого запису зі своїм же ID — не конфлікт
        resp = client.post(f'/api/records/{first.id}/edit', data=edit_form(first, EID))
        assert resp.status_code == 200


def test_operator_add_ignores_ehealth_id(app, client):
    with app.app_context():
        ensure_user('op', 'operator')
        make_record(ensure_user('op', 'operator').id)  # creates DeptA
        login(client, 'op')
        resp = client.post('/api/records/add', data={
            'date_of_discharge': DATE.isoformat(), 'full_name': 'Новий Пацієнт',
            'discharge_department': 'DeptA', 'treating_physician': 'Лікар',
            'history': '777', 'k_days': '3', 'patient_ehealth_id': EID,
        }, headers={'X-Requested-With': 'XMLHttpRequest'})
        assert resp.status_code == 200
        assert Record.query.filter_by(full_name='Новий Пацієнт').one().patient_ehealth_id is None


def test_db_enforces_uniqueness(app):
    from sqlalchemy.exc import IntegrityError
    with app.app_context():
        u = ensure_user('editor', 'editor')
        make_record(u.id, 'A', ehealth_id=EID)
        make_record(u.id, 'B')  # NULL-и не конфліктують
        make_record(u.id, 'C')
        with pytest.raises(IntegrityError):
            make_record(u.id, 'D', ehealth_id=EID)
        db.session.rollback()


# --- visibility, search, export ---

@pytest.mark.parametrize('role,visible', [('editor', True), ('admin', True),
                                          ('operator', False), ('viewer', False)])
def test_ehealth_id_visibility(app, client, role, visible):
    """Стовпця в таблиці немає ні в кого; значення доступне лише admin/editor (у формі редагування)."""
    with app.app_context():
        u = ensure_user(role, role)
        make_record(u.id, ehealth_id=EID)
        login(client, role)
        html = client.get('/', query_string={'all_months': '1'}).get_data(as_text=True)
        assert 'ID ЕСОЗ</th>' not in html
        assert (f'data-patient-ehealth-id="{EID}"' in html) is visible


@pytest.mark.parametrize('role,found', [('editor', True), ('operator', False)])
def test_search_by_ehealth_id(app, client, role, found):
    with app.app_context():
        u = ensure_user(role, role)
        target = make_record(u.id, 'Шуканий Пацієнт', ehealth_id=EID)
        other = make_record(u.id, 'Інший Пацієнт', ehealth_id=EID2)
        login(client, role)
        html = client.get('/', query_string={'all_months': '1', 'full_name': EID[:13]}).get_data(as_text=True)
        assert (f'id="record-{target.id}"' in html) is found
        assert f'id="record-{other.id}"' not in html


def test_export_has_ehealth_column(app, client):
    with app.app_context():
        u = ensure_user('editor', 'editor')
        make_record(u.id, ehealth_id=EID)
        login(client, 'editor')
        resp = client.post('/export', data={'export_mode': 'month', 'month_filter': DATE.strftime('%Y-%m')})
        assert resp.status_code == 200, resp.status_code
        ws = load_workbook(io.BytesIO(resp.data)).active
        headers = [c.value for c in ws[1]]
        col = headers.index('ID пацієнта (ЕСОЗ)')
        assert ws.cell(row=2, column=col + 1).value == EID


@pytest.mark.parametrize('role,visible', [('editor', True), ('admin', True),
                                          ('operator', False), ('viewer', False)])
def test_record_info_column_and_modal(app, client, role, visible):
    """Стовпець «Інфо» (значок у кожному рядку) і модалка — лише для admin/editor."""
    with app.app_context():
        u = ensure_user(role, role)
        make_record(u.id, 'З ідентифікатором', ehealth_id=EID)
        make_record(u.id, 'Без ідентифікатора')
        login(client, role)
        html = client.get('/', query_string={'all_months': '1'}).get_data(as_text=True)
        assert ('Інфо</th>' in html) is visible
        assert ('id="recordInfoModal"' in html) is visible
        # значок у кожному рядку — і з ID ЕСОЗ, і без нього
        assert html.count('record-info-btn"') == (2 if visible else 0)
