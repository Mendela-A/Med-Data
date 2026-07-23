# app/blueprints/admin/routes.py
"""
Admin routes
"""

from flask import render_template, redirect, url_for, flash, request, current_app, send_file
from flask_login import login_required, current_user
from datetime import datetime, date, timedelta
from io import BytesIO
from sqlalchemy import extract, case, func

from app.extensions import db, cache
from models import User, Department, Audit, Record, DailyReport, AmbulatoryRecord, NSZUCorrection, StatusOption, log_action
from decorators import role_required
from utils import clear_dropdown_cache, escape_like, get_distinct_audit_actions, clamp_per_page
from constants import (VALID_ROLES, STATUS_DISCHARGED, STATUS_PROCESSING, STATUS_VIOLATIONS,
                       STATUS_DECEASED, STATUS_NO_GROUP, STATUS_NO_EPISODE, TABS, DEFAULT_ROLE_TABS,
                       UKRAINIAN_MONTHS)
from . import admin_bp


# User Management Routes
@admin_bp.route('/users')
@role_required('admin')
def admin_users():
    users = User.query.order_by(User.username).all()
    return render_template('admin_users.html', users=users)


@admin_bp.route('/users/create', methods=['POST'])
@role_required('admin')
def admin_create_user():
    username = request.form.get('username', '').strip()
    password = request.form.get('password', '').strip()
    role = request.form.get('role', '').strip() or 'operator'

    if role not in VALID_ROLES:
        flash('Невірна роль користувача', 'warning')
        return redirect(url_for('admin.admin_users'))

    if not username or not password:
        flash('Ім\'я користувача та пароль обов\'язкові', 'warning')
        return redirect(url_for('admin.admin_users'))
    if len(password) < 8:
        flash('Пароль повинен містити щонайменше 8 символів', 'warning')
        return redirect(url_for('admin.admin_users'))
    if User.query.filter_by(username=username).first():
        flash('Ім\'я користувача вже зайнято', 'warning')
        return redirect(url_for('admin.admin_users'))

    u = User(username=username, role=role)
    u.set_password(password)
    db.session.add(u)
    db.session.flush()  # assigns u.id
    log_action(current_user.id, 'user.create', 'user', u.id, f'role={role}')
    db.session.commit()
    current_app.logger.info(f'User created: {username} by {current_user.username}')
    flash(f'Користувача {username} ({role}) успішно створено', 'success')
    return redirect(url_for('admin.admin_users'))


@admin_bp.route('/users/<int:user_id>/edit', methods=['GET', 'POST'])
@role_required('admin')
def admin_edit_user(user_id):
    u = db.get_or_404(User, user_id)

    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        role = request.form.get('role', '').strip()

        if not username:
            flash('Ім\'я користувача обов\'язкове', 'warning')
            return redirect(url_for('admin.admin_edit_user', user_id=user_id))

        # Check if username is taken by another user
        existing = User.query.filter_by(username=username).first()
        if existing and existing.id != user_id:
            flash('Ім\'я користувача вже зайнято', 'warning')
            return redirect(url_for('admin.admin_edit_user', user_id=user_id))

        if password and len(password) < 8:
            flash('Пароль повинен містити щонайменше 8 символів', 'warning')
            return redirect(url_for('admin.admin_edit_user', user_id=user_id))

        # Update username
        old_username = u.username
        u.username = username

        # Update password if provided
        if password:
            u.set_password(password)

        # Update role
        if role in VALID_ROLES:
            u.role = role

        # Update extra_permissions (tabs beyond role defaults) and revoked_permissions
        # (role-default tabs explicitly unchecked for this user)
        base_tabs = DEFAULT_ROLE_TABS.get(role, [])
        checked_tabs = {t for t in TABS if request.form.get(f'tab_{t}')}
        extra = [t for t in TABS if t not in base_tabs and t in checked_tabs]
        revoked = [t for t in TABS if t in base_tabs and t not in checked_tabs]
        u.extra_permissions = extra if extra else None
        u.revoked_permissions = revoked if revoked else None

        try:
            details = f'username={old_username}->{username}, role={role}'
            if revoked:
                details += f', revoked_tabs={revoked}'
            if extra:
                details += f', extra_tabs={extra}'
            if password:
                details += ', password_changed=True'
            log_action(current_user.id, 'user.update', 'user', u.id, details)
            db.session.commit()
        except Exception:
            db.session.rollback()
            current_app.logger.exception('Failed to update user')
            flash('Помилка при збереженні змін', 'danger')
            return redirect(url_for('admin.admin_edit_user', user_id=user_id))
        current_app.logger.info(f'User updated: {u.username} by {current_user.username}')
        flash(f'Користувача {u.username} успішно оновлено', 'success')
        return redirect(url_for('admin.admin_users'))

    return render_template('edit_user.html', user=u, tabs=TABS, default_role_tabs=DEFAULT_ROLE_TABS)


@admin_bp.route('/users/<int:user_id>/delete', methods=['POST'])
@role_required('admin')
def admin_delete_user(user_id):
    if current_user.id == user_id:
        flash('Ви не можете видалити самого себе', 'danger')
        return redirect(url_for('admin.admin_users'))
    u = db.get_or_404(User, user_id)
    saved_id = u.id
    saved_username = u.username
    try:
        db.session.delete(u)
        log_action(current_user.id, 'user.delete', 'user', saved_id, f'username={saved_username}')
        db.session.commit()
    except Exception:
        db.session.rollback()
        current_app.logger.exception('Failed to delete user')
        flash('Помилка при видаленні користувача', 'danger')
        return redirect(url_for('admin.admin_users'))
    current_app.logger.info(f'User deleted: {saved_username} by {current_user.username}')
    flash(f'Користувача {saved_username} видалено', 'danger')
    return redirect(url_for('admin.admin_users'))


# Department Management Routes
@admin_bp.route('/departments')
@role_required('admin')
def admin_departments():
    departments = Department.query.order_by(Department.row_no.nullslast(), Department.name).all()
    return render_template('admin_departments.html', departments=departments)


@admin_bp.route('/departments/create', methods=['POST'])
@role_required('admin')
def admin_create_department():
    name = request.form.get('name', '').strip()
    bed_profile_name = request.form.get('bed_profile_name', '').strip() or None
    row_no_str = request.form.get('row_no', '').strip()
    bed_capacity_str = request.form.get('bed_capacity', '').strip()

    if not name:
        flash('Назва відділення обов\'язкова', 'warning')
        return redirect(url_for('admin.admin_departments'))
    if Department.query.filter_by(name=name).first():
        flash('Відділення з такою назвою вже існує', 'warning')
        return redirect(url_for('admin.admin_departments'))

    row_no = None
    if row_no_str:
        try:
            row_no = int(row_no_str)
        except ValueError:
            flash('№ рядка має бути числом', 'warning')
            return redirect(url_for('admin.admin_departments'))

    bed_capacity = None
    if bed_capacity_str:
        try:
            bed_capacity = int(bed_capacity_str)
        except ValueError:
            flash('Ліжковий фонд має бути числом', 'warning')
            return redirect(url_for('admin.admin_departments'))

    d = Department(
        name=name,
        bed_profile_name=bed_profile_name,
        row_no=row_no,
        bed_capacity=bed_capacity
    )
    db.session.add(d)
    db.session.flush()  # assigns d.id
    log_action(current_user.id, 'department.create', 'department', d.id, f'name={name}, profile={bed_profile_name}, row={row_no}, beds={bed_capacity}')
    db.session.commit()
    # Clear dropdown cache after creating department
    clear_dropdown_cache()
    current_app.logger.info(f'Department created: {name} by {current_user.username}')
    flash(f'Відділення "{name}" успішно створено', 'success')
    return redirect(url_for('admin.admin_departments'))


@admin_bp.route('/departments/<int:dept_id>/edit', methods=['GET', 'POST'])
@role_required('admin')
def admin_edit_department(dept_id):
    d = db.get_or_404(Department, dept_id)

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        bed_profile_name = request.form.get('bed_profile_name', '').strip() or None
        row_no_str = request.form.get('row_no', '').strip()
        bed_capacity_str = request.form.get('bed_capacity', '').strip()

        if not name:
            flash('Назва відділення обов\'язкова', 'warning')
            return redirect(url_for('admin.admin_edit_department', dept_id=dept_id))

        existing = Department.query.filter_by(name=name).first()
        if existing and existing.id != dept_id:
            flash('Відділення з такою назвою вже існує', 'warning')
            return redirect(url_for('admin.admin_edit_department', dept_id=dept_id))

        row_no = None
        if row_no_str:
            try:
                row_no = int(row_no_str)
            except ValueError:
                flash('№ рядка має бути числом', 'warning')
                return redirect(url_for('admin.admin_edit_department', dept_id=dept_id))

        bed_capacity = None
        if bed_capacity_str:
            try:
                bed_capacity = int(bed_capacity_str)
            except ValueError:
                flash('Ліжковий фонд має бути числом', 'warning')
                return redirect(url_for('admin.admin_edit_department', dept_id=dept_id))

        old_name = d.name
        d.name = name
        d.bed_profile_name = bed_profile_name
        d.row_no = row_no
        d.bed_capacity = bed_capacity

        try:
            log_action(current_user.id, 'department.update', 'department', d.id, f'name={old_name}->{name}, profile={bed_profile_name}, row={row_no}, beds={bed_capacity}')
            db.session.commit()
            clear_dropdown_cache()
        except Exception:
            db.session.rollback()
            current_app.logger.exception('Failed to update department')
            flash('Помилка при збереженні змін', 'danger')
            return redirect(url_for('admin.admin_edit_department', dept_id=dept_id))

        current_app.logger.info(f'Department updated: {d.name} by {current_user.username}')
        flash(f'Відділення "{d.name}" успішно оновлено', 'success')
        return redirect(url_for('admin.admin_departments'))

    return render_template('edit_department.html', department=d)


@admin_bp.route('/departments/<int:dept_id>/delete', methods=['POST'])
@role_required('admin')
def admin_delete_department(dept_id):
    d = db.get_or_404(Department, dept_id)
    # prevent deletion if department in use
    in_use = Record.query.filter(Record.discharge_department == d.name).count()
    if in_use:
        flash(f'Неможливо видалити відділення "{d.name}" - використовується в {in_use} записах', 'danger')
        return redirect(url_for('admin.admin_departments'))
    saved_id = d.id
    saved_name = d.name
    db.session.delete(d)
    log_action(current_user.id, 'department.delete', 'department', saved_id, f'name={saved_name}')
    db.session.commit()
    # Clear dropdown cache after deleting department
    clear_dropdown_cache()
    current_app.logger.info(f'Department deleted: {saved_name} by {current_user.username}')
    flash(f'Відділення "{saved_name}" видалено', 'danger')
    return redirect(url_for('admin.admin_departments'))


# Status Dictionary Routes (scopes: ambulatory / records / nszu)
STATUS_COLORS = ('primary', 'success', 'info', 'warning', 'danger', 'secondary', 'dark')

STATUS_SCOPES = {
    'ambulatory': {'label': 'Амбулаторія', 'model': AmbulatoryRecord, 'column': 'discharge_status'},
    'records': {'label': 'Записи (стаціонар)', 'model': Record, 'column': 'discharge_status'},
    'nszu': {'label': 'НСЗУ', 'model': NSZUCorrection, 'column': 'status'},
}


def _valid_scope(scope):
    return scope if scope in STATUS_SCOPES else 'ambulatory'


def _scope_status_column(scope):
    cfg = STATUS_SCOPES[scope]
    return getattr(cfg['model'], cfg['column'])


def _status_usage_counts(scope):
    """Кількість записів відповідного розділу на кожен статус (одним GROUP BY)."""
    col = _scope_status_column(scope)
    model = STATUS_SCOPES[scope]['model']
    rows = db.session.query(col, func.count(model.id)).group_by(col).all()
    return {name: cnt for name, cnt in rows if name}


@admin_bp.route('/statuses')
@role_required('admin')
def admin_statuses():
    scope = _valid_scope(request.args.get('scope', 'ambulatory'))
    statuses = (StatusOption.query.filter_by(scope=scope)
                .order_by(StatusOption.sort_order, StatusOption.name).all())
    usage = _status_usage_counts(scope)
    known = {s.name for s in statuses}
    orphans = {name: cnt for name, cnt in usage.items() if name not in known}
    return render_template('admin_statuses.html',
                           statuses=statuses, usage=usage, orphans=orphans,
                           colors=STATUS_COLORS, scope=scope,
                           scopes={k: v['label'] for k, v in STATUS_SCOPES.items()})


@admin_bp.route('/statuses/create', methods=['POST'])
@role_required('admin')
def admin_create_status():
    scope = _valid_scope(request.form.get('scope', 'ambulatory'))
    name = request.form.get('name', '').strip()
    color = request.form.get('color', '').strip()
    icon = request.form.get('icon', '').strip() or 'bi-circle'
    show_in_stats = request.form.get('show_in_stats') == 'on'

    if not name:
        flash('Назва статусу обов\'язкова', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=scope))
    if color not in STATUS_COLORS:
        color = 'secondary'
    if StatusOption.query.filter_by(scope=scope, name=name).first():
        flash('Статус з такою назвою вже існує', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=scope))

    max_order = db.session.query(func.max(StatusOption.sort_order)).filter_by(scope=scope).scalar() or 0
    s = StatusOption(scope=scope, name=name, color=color, icon=icon,
                     sort_order=max_order + 10, show_in_stats=show_in_stats)
    db.session.add(s)
    db.session.flush()
    log_action(current_user.id, 'status.create', 'status_option', s.id, f'scope={scope}, name={name}')
    db.session.commit()
    clear_dropdown_cache()
    current_app.logger.info(f'StatusOption created: [{scope}] {name} by {current_user.username}')
    flash(f'Статус «{name}» успішно створено', 'success')
    return redirect(url_for('admin.admin_statuses', scope=scope))


@admin_bp.route('/statuses/<int:status_id>/update', methods=['POST'])
@role_required('admin')
def admin_update_status(status_id):
    s = db.get_or_404(StatusOption, status_id)
    new_name = request.form.get('name', '').strip()
    color = request.form.get('color', '').strip()
    icon = request.form.get('icon', '').strip() or s.icon
    sort_order = request.form.get('sort_order', type=int)
    show_in_stats = request.form.get('show_in_stats') == 'on'

    if not new_name:
        flash('Назва статусу обов\'язкова', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))

    old_name = s.name
    renamed = new_name != old_name
    if renamed and s.is_system:
        flash(f'Статус «{old_name}» — системний (на ньому тримаються статистика і звіти), його не можна перейменувати', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))
    if renamed and StatusOption.query.filter_by(scope=s.scope, name=new_name).first():
        flash('Статус з такою назвою вже існує', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))

    s.name = new_name
    if color in STATUS_COLORS:
        s.color = color
    s.icon = icon
    if sort_order is not None:
        s.sort_order = sort_order
    s.show_in_stats = show_in_stats

    try:
        renamed_count = 0
        if renamed:
            # Записи зберігають статус текстом — перейменування мусить
            # оновити їх в тій самій транзакції, інакше фільтри/піли
            # «загублять» старі записи.
            col = _scope_status_column(s.scope)
            model = STATUS_SCOPES[s.scope]['model']
            renamed_count = (model.query
                             .filter(col == old_name)
                             .update({STATUS_SCOPES[s.scope]['column']: new_name},
                                     synchronize_session=False))
        details = f'scope={s.scope}, name={old_name}->{new_name}, color={s.color}, records_renamed={renamed_count}'
        log_action(current_user.id, 'status.update', 'status_option', s.id, details)
        db.session.commit()
    except Exception:
        db.session.rollback()
        current_app.logger.exception('Failed to update status option')
        flash('Помилка при збереженні статусу', 'danger')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))

    clear_dropdown_cache()
    current_app.logger.info(f'StatusOption updated: [{s.scope}] {old_name}->{new_name} by {current_user.username}')
    if renamed and renamed_count:
        flash(f'Статус «{old_name}» перейменовано на «{new_name}», оновлено записів: {renamed_count}', 'success')
    else:
        flash(f'Статус «{new_name}» оновлено', 'success')
    return redirect(url_for('admin.admin_statuses', scope=s.scope))


@admin_bp.route('/statuses/<int:status_id>/set-default', methods=['POST'])
@role_required('admin')
def admin_set_default_status(status_id):
    s = db.get_or_404(StatusOption, status_id)
    if not s.is_active:
        flash('Неактивний статус не може бути статусом за замовчуванням', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))
    StatusOption.query.filter_by(scope=s.scope).update({'is_default': False}, synchronize_session=False)
    s.is_default = True
    log_action(current_user.id, 'status.set_default', 'status_option', s.id, f'scope={s.scope}, name={s.name}')
    db.session.commit()
    clear_dropdown_cache()
    flash(f'Статус «{s.name}» встановлено за замовчуванням для нових записів', 'success')
    return redirect(url_for('admin.admin_statuses', scope=s.scope))


@admin_bp.route('/statuses/<int:status_id>/toggle', methods=['POST'])
@role_required('admin')
def admin_toggle_status(status_id):
    s = db.get_or_404(StatusOption, status_id)
    if s.is_system:
        flash(f'Статус «{s.name}» — системний, його не можна деактивувати', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))
    if s.is_active and s.is_default:
        flash('Статус за замовчуванням не можна деактивувати — спочатку призначте інший', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))
    s.is_active = not s.is_active
    action = 'status.activate' if s.is_active else 'status.deactivate'
    log_action(current_user.id, action, 'status_option', s.id, f'scope={s.scope}, name={s.name}')
    db.session.commit()
    clear_dropdown_cache()
    state = 'активовано' if s.is_active else 'деактивовано'
    flash(f'Статус «{s.name}» {state}', 'success')
    return redirect(url_for('admin.admin_statuses', scope=s.scope))


@admin_bp.route('/statuses/<int:status_id>/delete', methods=['POST'])
@role_required('admin')
def admin_delete_status(status_id):
    s = db.get_or_404(StatusOption, status_id)
    if s.is_system:
        flash(f'Статус «{s.name}» — системний, його не можна видалити', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))
    col = _scope_status_column(s.scope)
    in_use = STATUS_SCOPES[s.scope]['model'].query.filter(col == s.name).count()
    if in_use:
        flash(f'Неможливо видалити статус «{s.name}» — використовується в {in_use} записах. Деактивуйте його натомість.', 'danger')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))
    if s.is_default:
        flash('Статус за замовчуванням не можна видалити — спочатку призначте інший', 'warning')
        return redirect(url_for('admin.admin_statuses', scope=s.scope))
    saved_id, saved_name, saved_scope = s.id, s.name, s.scope
    db.session.delete(s)
    log_action(current_user.id, 'status.delete', 'status_option', saved_id, f'scope={saved_scope}, name={saved_name}')
    db.session.commit()
    clear_dropdown_cache()
    current_app.logger.info(f'StatusOption deleted: [{saved_scope}] {saved_name} by {current_user.username}')
    flash(f'Статус «{saved_name}» видалено', 'danger')
    return redirect(url_for('admin.admin_statuses', scope=saved_scope))


# Statistics Route
@admin_bp.route('/statistics')
@role_required('admin', 'viewer')
def admin_statistics():
    now = datetime.now()
    today = now.date()

    # --- Parse date range from query params ---
    from_str = request.args.get('from_date', '').strip()
    to_str = request.args.get('to_date', '').strip()

    # Backward compatibility: month_year param (YYYY-MM)
    month_year = request.args.get('month_year', '')

    from_date = None
    to_date = None

    # Try from_date / to_date first
    if from_str:
        try:
            from_date = date.fromisoformat(from_str)
        except ValueError:
            pass
    if to_str:
        try:
            to_date = date.fromisoformat(to_str)
        except ValueError:
            pass

    # Fallback: month_year param
    if from_date is None and month_year and '-' in month_year:
        try:
            y, m = month_year.split('-')
            y, m = int(y), int(m)
            if 1 <= m <= 12 and 2000 <= y <= 2100:
                from_date = date(y, m, 1)
                to_date = date(y + 1, 1, 1) - timedelta(days=1) if m == 12 else date(y, m + 1, 1) - timedelta(days=1)
        except (ValueError, IndexError):
            pass

    # Default: current month
    if from_date is None:
        from_date = date(today.year, today.month, 1)
    if to_date is None:
        if from_date.month == 12:
            to_date = date(from_date.year + 1, 1, 1) - timedelta(days=1)
        else:
            to_date = date(from_date.year, from_date.month + 1, 1) - timedelta(days=1)

    # Validate: from <= to
    if from_date > to_date:
        from_date, to_date = to_date, from_date

    # Exclusive upper bound for queries (to_date is inclusive, so +1 day)
    query_end = to_date + timedelta(days=1)

    # Period label for display (українські назви місяців, не залежимо від локалі)
    if from_date.day == 1 and to_date == (date(from_date.year, from_date.month + 1, 1) - timedelta(days=1) if from_date.month < 12 else date(from_date.year + 1, 1, 1) - timedelta(days=1)):
        period_label = f"{UKRAINIAN_MONTHS[from_date.month]} {from_date.year}"
    else:
        period_label = f"{from_date.strftime('%d.%m.%Y')} — {to_date.strftime('%d.%m.%Y')}"

    # Cache key — 5-minute TTL, no manual invalidation needed
    _cache_key = f'stats_{from_date}_{to_date}'
    _cached = cache.get(_cache_key)
    if _cached:
        return render_template('admin_statistics.html',
                               period_label=period_label,
                               from_date=from_date, to_date=to_date,
                               **_cached)

    # 1. Records per day by discharge date
    per_day_rows = db.session.query(
        func.date(Record.date_of_discharge).label('date'),
        func.count(Record.id).label('count')
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end
    ).group_by(
        Record.date_of_discharge
    ).order_by(Record.date_of_discharge).all()

    max_per_day = max((r.count for r in per_day_rows), default=0)
    records_per_day = []
    for r in per_day_rows:
        d = r.date if not isinstance(r.date, str) else datetime.strptime(r.date, '%Y-%m-%d').date()
        records_per_day.append({'date': d.strftime('%d.%m.%Y'), 'count': r.count})

    # 2. Status distribution by department + ALOS
    dept_stats = db.session.query(
        Record.discharge_department,
        func.sum(case((Record.date_of_death.isnot(None), 1), else_=0)).label('deceased'),
        func.sum(case(((Record.discharge_status == STATUS_DISCHARGED) & (Record.date_of_death.is_(None)), 1), else_=0)).label('discharged'),
        func.sum(case(((Record.discharge_status == STATUS_PROCESSING) & (Record.date_of_death.is_(None)), 1), else_=0)).label('processing'),
        func.sum(case(((Record.discharge_status == STATUS_VIOLATIONS) & (Record.date_of_death.is_(None)), 1), else_=0)).label('violations'),
        func.avg(Record.k_days).label('avg_k_days')
    ).filter(
        Record.discharge_department.isnot(None),
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end
    ).group_by(Record.discharge_department).all()

    status_by_dept = {}
    for row in dept_stats:
        dept, deceased, discharged, processing, violations, avg_k = (
            row.discharge_department, row.deceased, row.discharged,
            row.processing, row.violations, row.avg_k_days)
        status_by_dept[dept] = {
            STATUS_DECEASED:   deceased or 0,
            STATUS_DISCHARGED: discharged or 0,
            STATUS_PROCESSING: processing or 0,
            STATUS_VIOLATIONS: violations or 0,
            'avg_k_days': round(avg_k, 1) if avg_k is not None else None,
        }

    dept_list = sorted(status_by_dept.keys())

    # H3: compute status_distribution from status_by_dept (removes redundant current_stats query)
    status_distribution = {
        STATUS_DECEASED:   sum(d[STATUS_DECEASED]   for d in status_by_dept.values()),
        STATUS_DISCHARGED: sum(d[STATUS_DISCHARGED] for d in status_by_dept.values()),
        STATUS_PROCESSING: sum(d[STATUS_PROCESSING] for d in status_by_dept.values()),
        STATUS_VIOLATIONS: sum(d[STATUS_VIOLATIONS] for d in status_by_dept.values()),
    }
    total_records = sum(status_distribution.values())

    # H2: global ALOS
    all_k = [d['avg_k_days'] for d in status_by_dept.values() if d['avg_k_days'] is not None]
    global_alos = round(sum(all_k) / len(all_k), 1) if all_k else None

    # АДСЖ group breakdown
    adsj_raw = db.session.query(
        func.coalesce(
            func.nullif(func.trim(Record.adsj), ''),
            STATUS_NO_GROUP
        ).label('group_name'),
        func.count(Record.id).label('count'),
        func.sum(Record.suma).label('total_suma')
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end
    ).group_by(
        func.coalesce(
            func.nullif(func.trim(Record.adsj), ''),
            STATUS_NO_GROUP
        )
    ).all()

    adsj_stats = sorted(
        [r for r in adsj_raw if r.group_name != STATUS_NO_GROUP],
        key=lambda r: r.group_name
    ) + [r for r in adsj_raw if r.group_name == STATUS_NO_GROUP]

    adsj_total_count = sum(r.count for r in adsj_stats)
    adsj_total_suma = sum(r.total_suma or 0 for r in adsj_stats)

    # --- Trends: compare with previous period of equal length ---
    range_days = (to_date - from_date).days + 1
    prev_to = from_date - timedelta(days=1)
    prev_from = prev_to - timedelta(days=range_days - 1)
    prev_query_end = prev_to + timedelta(days=1)

    prev_stats = db.session.query(
        func.sum(case((Record.date_of_death.isnot(None), 1), else_=0)).label('deceased'),
        func.sum(case(((Record.discharge_status == STATUS_DISCHARGED) & (Record.date_of_death.is_(None)), 1), else_=0)).label('discharged'),
        func.sum(case(((Record.discharge_status == STATUS_PROCESSING) & (Record.date_of_death.is_(None)), 1), else_=0)).label('processing'),
        func.sum(case(((Record.discharge_status == STATUS_VIOLATIONS) & (Record.date_of_death.is_(None)), 1), else_=0)).label('violations')
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= prev_from,
        Record.date_of_discharge < prev_query_end
    ).first()

    prev_deceased   = prev_stats.deceased or 0
    prev_discharged = prev_stats.discharged or 0
    prev_processing = prev_stats.processing or 0
    prev_violations = prev_stats.violations or 0
    prev_total = prev_deceased + prev_discharged + prev_processing + prev_violations

    trends = {
        'total':      total_records - prev_total,
        'processing': status_distribution[STATUS_PROCESSING] - prev_processing,
        'discharged': status_distribution[STATUS_DISCHARGED] - prev_discharged,
        'deceased':   status_distribution[STATUS_DECEASED]   - prev_deceased,
        'violations': status_distribution[STATUS_VIOLATIONS] - prev_violations,
    }

    # H1: Bed occupancy from DailyReport
    bed_rows = db.session.query(
        Department.name,
        func.avg(DailyReport.beds_total).label('avg_beds'),
        func.avg(DailyReport.patients_end).label('avg_occupied'),
        func.sum(DailyReport.admitted_total).label('total_admitted'),
        func.sum(DailyReport.deaths).label('total_deaths'),
    ).join(Department, DailyReport.department_id == Department.id
    ).filter(
        DailyReport.report_date >= from_date,
        DailyReport.report_date < query_end,
        DailyReport.beds_total.isnot(None),
    ).group_by(DailyReport.department_id, Department.name).all()

    bed_stats = sorted([{
        'name':         r.name,
        'avg_beds':     round(r.avg_beds or 0),
        'avg_occupied': round(r.avg_occupied or 0),
        'occupancy_pct': round((r.avg_occupied or 0) / max(r.avg_beds or 1, 1) * 100),
        'admitted':     r.total_admitted or 0,
        'deaths':       r.total_deaths or 0,
    } for r in bed_rows], key=lambda x: x['name'])

    global_bed = {
        'avg_beds':      sum(d['avg_beds'] for d in bed_stats),
        'avg_occupied':  sum(d['avg_occupied'] for d in bed_stats),
        'occupancy_pct': round(
            sum(d['avg_occupied'] for d in bed_stats) /
            max(sum(d['avg_beds'] for d in bed_stats), 1) * 100
        ) if bed_stats else 0,
        'total_admitted': sum(d['admitted'] for d in bed_stats),
        'total_deaths':   sum(d['deaths'] for d in bed_stats),
    }

    # H1: Ambulatory record stats
    amb_r = db.session.query(
        func.count(AmbulatoryRecord.id).label('total'),
        func.sum(case((AmbulatoryRecord.is_urgent == True, 1), else_=0)).label('urgent'),
        func.sum(case((AmbulatoryRecord.discharge_status == STATUS_DISCHARGED, 1), else_=0)).label('discharged'),
        func.sum(case((AmbulatoryRecord.discharge_status == STATUS_NO_EPISODE, 1), else_=0)).label('no_episode'),
    ).filter(
        AmbulatoryRecord.date >= from_date,
        AmbulatoryRecord.date < query_end,
    ).first()
    amb_stats = {
        'total':      amb_r.total or 0,
        'urgent':     amb_r.urgent or 0,
        'discharged': amb_r.discharged or 0,
        'no_episode': amb_r.no_episode or 0,
    }

    ctx = dict(
        records_per_day=records_per_day,
        max_per_day=max_per_day,
        status_by_dept=status_by_dept,
        dept_list=dept_list,
        status_distribution=status_distribution,
        total_records=total_records,
        global_alos=global_alos,
        trends=trends,
        adsj_stats=adsj_stats,
        adsj_total_count=adsj_total_count,
        adsj_total_suma=adsj_total_suma,
        bed_stats=bed_stats,
        global_bed=global_bed,
        amb_stats=amb_stats,
    )
    cache.set(_cache_key, ctx, timeout=300)

    return render_template(
        'admin_statistics.html',
        period_label=period_label,
        from_date=from_date,
        to_date=to_date,
        **ctx,
    )


_SUBMISSION_EXCL_DEPTS = ['гінекологія', 'реанімація']


def _parse_report_dates():
    today = datetime.now().date()
    from_str = request.args.get('from_date', '').strip()
    to_str = request.args.get('to_date', '').strip()
    from_date = None
    to_date = None
    if from_str:
        try:
            from_date = date.fromisoformat(from_str)
        except ValueError:
            pass
    if to_str:
        try:
            to_date = date.fromisoformat(to_str)
        except ValueError:
            pass
    if from_date is None:
        from_date = date(today.year, today.month, 1)
    if to_date is None:
        if from_date.month == 12:
            to_date = date(from_date.year + 1, 1, 1) - timedelta(days=1)
        else:
            to_date = date(from_date.year, from_date.month + 1, 1) - timedelta(days=1)
    if from_date > to_date:
        from_date, to_date = to_date, from_date
    query_end = to_date + timedelta(days=1)
    if from_date.month == to_date.month and from_date.year == to_date.year and from_date.day == 1:
        period_label = f"{UKRAINIAN_MONTHS[from_date.month]} {from_date.year}"
    else:
        period_label = f"{from_date.strftime('%d.%m.%Y')} — {to_date.strftime('%d.%m.%Y')}"
    return from_date, to_date, query_end, period_label


def _physician_records(physician, from_date, query_end):
    """Записи одного лікаря за період (з тими ж виключеннями, що й звіт)."""
    return Record.query.filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        Record.treating_physician == physician,
        func.lower(Record.discharge_department).notin_(_SUBMISSION_EXCL_DEPTS),
    ).order_by(
        Record.history_submitted.asc(),            # спершу «не здано»
        Record.date_of_discharge.desc(),
    ).all()


# Reports Route
@admin_bp.route('/reports')
@role_required('operator', 'editor', 'admin', 'viewer')
def admin_reports():
    return redirect(url_for('admin.report_submission_page'))


@admin_bp.route('/reports/submission-page')
@role_required('operator', 'editor', 'admin', 'viewer')
def report_submission_page():
    from_date, to_date, query_end, period_label = _parse_report_dates()

    submission_row = db.session.query(
        func.sum(case((Record.history_submitted == True, 1), else_=0)).label('submitted'),
        func.sum(case((Record.history_submitted == False, 1), else_=0)).label('not_submitted'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        func.lower(Record.discharge_department).notin_(_SUBMISSION_EXCL_DEPTS),
    ).first()

    submission_by_physician = db.session.query(
        Record.treating_physician,
        func.sum(case((Record.history_submitted == True, 1), else_=0)).label('submitted'),
        func.sum(case((Record.history_submitted == False, 1), else_=0)).label('not_submitted'),
        func.count(Record.id).label('total'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        Record.treating_physician.isnot(None),
        func.trim(Record.treating_physician) != '',
        func.lower(Record.discharge_department).notin_(_SUBMISSION_EXCL_DEPTS),
    ).group_by(Record.treating_physician).order_by(func.count(Record.id).desc()).all()

    return render_template(
        'report_submission.html',
        from_date=from_date,
        to_date=to_date,
        period_label=period_label,
        submission_submitted=submission_row.submitted or 0,
        submission_not_submitted=submission_row.not_submitted or 0,
        submission_by_physician=submission_by_physician,
    )


@admin_bp.route('/reports/urgency-page')
@role_required('operator', 'editor', 'admin', 'viewer')
def report_urgency_page():
    from_date, to_date, query_end, period_label = _parse_report_dates()

    urgency_row = db.session.query(
        func.sum(case((Record.is_urgent == True, 1), else_=0)).label('urgent'),
        func.sum(case((Record.is_urgent == False, 1), else_=0)).label('planned'),
        func.sum(case((Record.is_urgent.is_(None), 1), else_=0)).label('unset'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
    ).first()

    urgency_by_dept = db.session.query(
        Record.discharge_department,
        func.sum(case((Record.is_urgent == True, 1), else_=0)).label('urgent'),
        func.sum(case((Record.is_urgent == False, 1), else_=0)).label('planned'),
        func.sum(case((Record.is_urgent.is_(None), 1), else_=0)).label('unset'),
        func.count(Record.id).label('total'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        Record.discharge_department.isnot(None),
        func.trim(Record.discharge_department) != '',
        func.lower(Record.discharge_department).notin_(['гінекологічне', 'гінекологія']),
    ).group_by(Record.discharge_department).order_by(func.count(Record.id).desc()).all()

    return render_template(
        'report_urgency.html',
        from_date=from_date,
        to_date=to_date,
        period_label=period_label,
        urgency_urgent=urgency_row.urgent or 0,
        urgency_planned=urgency_row.planned or 0,
        urgency_unset=urgency_row.unset or 0,
        urgency_by_dept=urgency_by_dept,
    )



@admin_bp.route('/reports/submission')
@role_required('operator', 'editor', 'admin', 'viewer')
def report_submission():
    """PDF: history submission stats per physician."""
    from_date, to_date, query_end, _ = _parse_report_dates()

    submission_row = db.session.query(
        func.sum(case((Record.history_submitted == True, 1), else_=0)).label('submitted'),
        func.sum(case((Record.history_submitted == False, 1), else_=0)).label('not_submitted'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        func.lower(Record.discharge_department).notin_(_SUBMISSION_EXCL_DEPTS),
    ).first()

    submission_by_physician = db.session.query(
        Record.treating_physician,
        func.sum(case((Record.history_submitted == True, 1), else_=0)).label('submitted'),
        func.sum(case((Record.history_submitted == False, 1), else_=0)).label('not_submitted'),
        func.count(Record.id).label('total'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        Record.treating_physician.isnot(None),
        func.trim(Record.treating_physician) != '',
        func.lower(Record.discharge_department).notin_(_SUBMISSION_EXCL_DEPTS),
    ).group_by(Record.treating_physician).order_by(
        func.sum(case((Record.history_submitted == False, 1), else_=0)).desc()
    ).all()

    import datetime as _dt
    kyiv_tz = _dt.timezone(_dt.timedelta(hours=2))

    try:
        from weasyprint import HTML
    except ImportError:
        flash('Для формування PDF потрібен пакет WeasyPrint', 'danger')
        return redirect(url_for('admin.admin_reports'))

    html_string = render_template(
        'print_submission.html',
        from_date=from_date,
        to_date=to_date,
        filter_physician='',
        submission_submitted=submission_row.submitted or 0,
        submission_not_submitted=submission_row.not_submitted or 0,
        submission_by_physician=submission_by_physician,
        generated_by=current_user.username,
        generated_at=datetime.now(kyiv_tz),
    )
    pdf = HTML(string=html_string).write_pdf()
    bio = BytesIO(pdf)
    bio.seek(0)
    filename = f"submission_{from_date.strftime('%d-%m-%Y')}_{to_date.strftime('%d-%m-%Y')}.pdf"
    try:
        log_action(current_user.id, 'admin.report_submission', 'report', None,
                   f'from={from_date} to={to_date}')
        db.session.commit()
    except Exception:
        current_app.logger.exception('Failed to log report_submission')
    return send_file(bio, as_attachment=True, download_name=filename, mimetype='application/pdf')


_SUBMISSION_PHYSICIAN_COLS = ['№', 'Дата виписки', 'ПІБ', 'Відділення', '№ історії хвороби', 'Статус виписки']


def _filename_part(text):
    """Безпечний фрагмент імені файлу з імені лікаря (зберігає кирилицю)."""
    safe = ''.join(ch if ch.isalnum() else '_' for ch in (text or '').strip())
    return safe.strip('_') or 'physician'


@admin_bp.route('/reports/submission-page/physician')
@role_required('operator', 'editor', 'admin', 'viewer')
def report_submission_physician():
    """HTML: перелік історій одного лікаря (здано / не здано) за період."""
    physician = request.args.get('physician', '').strip()
    from_date, to_date, query_end, period_label = _parse_report_dates()

    if not physician:
        flash('Не вказано лікаря', 'warning')
        return redirect(url_for('admin.report_submission_page',
                                from_date=from_date.isoformat(), to_date=to_date.isoformat()))

    records = _physician_records(physician, from_date, query_end)
    not_submitted_records = [r for r in records if not r.history_submitted]
    submitted_records = [r for r in records if r.history_submitted]

    return render_template(
        'report_submission_physician.html',
        physician=physician,
        from_date=from_date,
        to_date=to_date,
        period_label=period_label,
        not_submitted_records=not_submitted_records,
        submitted_records=submitted_records,
    )


@admin_bp.route('/reports/submission/physician.pdf')
@role_required('operator', 'editor', 'admin', 'viewer')
def report_submission_physician_pdf():
    """PDF: перелік історій одного лікаря (здано / не здано) за період."""
    physician = request.args.get('physician', '').strip()
    from_date, to_date, query_end, _ = _parse_report_dates()

    if not physician:
        flash('Не вказано лікаря', 'warning')
        return redirect(url_for('admin.report_submission_page',
                                from_date=from_date.isoformat(), to_date=to_date.isoformat()))

    records = _physician_records(physician, from_date, query_end)
    not_submitted_records = [r for r in records if not r.history_submitted]
    submitted_records = [r for r in records if r.history_submitted]

    try:
        from weasyprint import HTML
    except ImportError:
        flash('Для формування PDF потрібен пакет WeasyPrint', 'danger')
        return redirect(url_for('admin.report_submission_physician',
                                physician=physician,
                                from_date=from_date.isoformat(), to_date=to_date.isoformat()))

    import datetime as _dt
    kyiv_tz = _dt.timezone(_dt.timedelta(hours=2))

    html_string = render_template(
        'print_submission_physician.html',
        physician=physician,
        from_date=from_date,
        to_date=to_date,
        not_submitted_records=not_submitted_records,
        submitted_records=submitted_records,
        generated_by=current_user.username,
        generated_at=datetime.now(kyiv_tz),
    )
    pdf = HTML(string=html_string).write_pdf()
    bio = BytesIO(pdf)
    bio.seek(0)
    filename = (f"submission_{_filename_part(physician)}_"
                f"{from_date.strftime('%d-%m-%Y')}_{to_date.strftime('%d-%m-%Y')}.pdf")
    try:
        log_action(current_user.id, 'admin.report_submission_physician', 'report', None,
                   f'physician={physician} from={from_date} to={to_date}')
        db.session.commit()
    except Exception:
        current_app.logger.exception('Failed to log report_submission_physician')
    return send_file(bio, as_attachment=True, download_name=filename, mimetype='application/pdf')


def _style_xlsx_header(ws):
    """Стиль рядка заголовків (як у records.export)."""
    from openpyxl.styles import Font, PatternFill, Alignment
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')


def _autosize_xlsx(ws):
    """Авторозмір колонок (як у records.export)."""
    from openpyxl.utils import get_column_letter
    for i, col in enumerate(ws.columns, 1):
        max_length = 0
        column = get_column_letter(i)
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)


@admin_bp.route('/reports/submission/physician.xlsx')
@role_required('operator', 'editor', 'admin', 'viewer')
def report_submission_physician_xlsx():
    """Excel: перелік історій одного лікаря (здано / не здано) за період."""
    from openpyxl import Workbook

    physician = request.args.get('physician', '').strip()
    from_date, to_date, query_end, _ = _parse_report_dates()

    if not physician:
        flash('Не вказано лікаря', 'warning')
        return redirect(url_for('admin.report_submission_page',
                                from_date=from_date.isoformat(), to_date=to_date.isoformat()))

    records = _physician_records(physician, from_date, query_end)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Здача'
    ws.append(_SUBMISSION_PHYSICIAN_COLS + ['Здано'])
    for idx, r in enumerate(records, 1):
        ws.append([
            idx,
            r.date_of_discharge.strftime('%d.%m.%Y') if r.date_of_discharge else '',
            r.full_name,
            r.discharge_department or '',
            r.history or '',
            r.discharge_status or '',
            'Так' if r.history_submitted else 'Ні',
        ])
    _style_xlsx_header(ws)
    _autosize_xlsx(ws)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = (f"submission_{_filename_part(physician)}_"
                f"{from_date.strftime('%d-%m-%Y')}_{to_date.strftime('%d-%m-%Y')}.xlsx")
    try:
        log_action(current_user.id, 'admin.report_submission_physician', 'export', None,
                   f'physician={physician} from={from_date} to={to_date} count={len(records)}')
        db.session.commit()
    except Exception:
        current_app.logger.exception('Failed to log report_submission_physician xlsx')
    return send_file(bio, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@admin_bp.route('/reports/submission.xlsx')
@role_required('operator', 'editor', 'admin', 'viewer')
def report_submission_xlsx():
    """Excel «по всіх лікарях»: лист «Зведення» + лист «Всі записи»."""
    from openpyxl import Workbook

    from_date, to_date, query_end, _ = _parse_report_dates()

    submission_by_physician = db.session.query(
        Record.treating_physician,
        func.sum(case((Record.history_submitted == True, 1), else_=0)).label('submitted'),
        func.sum(case((Record.history_submitted == False, 1), else_=0)).label('not_submitted'),
        func.count(Record.id).label('total'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        Record.treating_physician.isnot(None),
        func.trim(Record.treating_physician) != '',
        func.lower(Record.discharge_department).notin_(_SUBMISSION_EXCL_DEPTS),
    ).group_by(Record.treating_physician).order_by(func.count(Record.id).desc()).all()

    all_records = Record.query.filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        Record.treating_physician.isnot(None),
        func.trim(Record.treating_physician) != '',
        func.lower(Record.discharge_department).notin_(_SUBMISSION_EXCL_DEPTS),
    ).order_by(
        func.lower(Record.treating_physician).asc(),
        Record.history_submitted.asc(),
        Record.date_of_discharge.desc(),
    ).all()

    wb = Workbook()

    # Лист 1 — Зведення
    ws_sum = wb.active
    ws_sum.title = 'Зведення'
    ws_sum.append(['Лікар', 'Здано', 'Не здано', 'Всього', '% здачі'])
    tot_submitted = tot_not = tot_all = 0
    for row in submission_by_physician:
        pct = round(row.submitted / row.total * 100) if row.total else 0
        ws_sum.append([row.treating_physician, row.submitted, row.not_submitted, row.total, f'{pct}%'])
        tot_submitted += row.submitted
        tot_not += row.not_submitted
        tot_all += row.total
    tot_pct = round(tot_submitted / tot_all * 100) if tot_all else 0
    ws_sum.append(['Всього', tot_submitted, tot_not, tot_all, f'{tot_pct}%'])
    _style_xlsx_header(ws_sum)
    _autosize_xlsx(ws_sum)

    # Лист 2 — Всі записи
    ws_det = wb.create_sheet('Всі записи')
    ws_det.append(['№', 'Дата виписки', 'ПІБ', 'Відділення', 'Лікар', '№ історії хвороби', 'Статус виписки', 'Здано'])
    for idx, r in enumerate(all_records, 1):
        ws_det.append([
            idx,
            r.date_of_discharge.strftime('%d.%m.%Y') if r.date_of_discharge else '',
            r.full_name,
            r.discharge_department or '',
            r.treating_physician,
            r.history or '',
            r.discharge_status or '',
            'Так' if r.history_submitted else 'Ні',
        ])
    _style_xlsx_header(ws_det)
    _autosize_xlsx(ws_det)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"submission_all_{from_date.strftime('%d-%m-%Y')}_{to_date.strftime('%d-%m-%Y')}.xlsx"
    try:
        log_action(current_user.id, 'admin.report_submission_all', 'export', None,
                   f'from={from_date} to={to_date} count={len(all_records)}')
        db.session.commit()
    except Exception:
        current_app.logger.exception('Failed to log report_submission_all xlsx')
    return send_file(bio, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@admin_bp.route('/reports/urgency')
@role_required('operator', 'editor', 'admin', 'viewer')
def report_urgency():
    """PDF: urgency stats per department."""
    from_date, to_date, query_end, _ = _parse_report_dates()

    urgency_row = db.session.query(
        func.sum(case((Record.is_urgent == True, 1), else_=0)).label('urgent'),
        func.sum(case((Record.is_urgent == False, 1), else_=0)).label('planned'),
        func.sum(case((Record.is_urgent.is_(None), 1), else_=0)).label('unset'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
    ).first()

    urgency_by_dept = db.session.query(
        Record.discharge_department,
        func.sum(case((Record.is_urgent == True, 1), else_=0)).label('urgent'),
        func.sum(case((Record.is_urgent == False, 1), else_=0)).label('planned'),
        func.sum(case((Record.is_urgent.is_(None), 1), else_=0)).label('unset'),
        func.count(Record.id).label('total'),
    ).filter(
        Record.date_of_discharge.isnot(None),
        Record.date_of_discharge >= from_date,
        Record.date_of_discharge < query_end,
        Record.discharge_department.isnot(None),
        func.trim(Record.discharge_department) != '',
        func.lower(Record.discharge_department).notin_(['гінекологічне', 'гінекологія']),
    ).group_by(Record.discharge_department).order_by(
        func.sum(case((Record.is_urgent == True, 1), else_=0)).desc()
    ).all()

    import datetime as _dt
    kyiv_tz = _dt.timezone(_dt.timedelta(hours=2))

    try:
        from weasyprint import HTML
    except ImportError:
        flash('Для формування PDF потрібен пакет WeasyPrint', 'danger')
        return redirect(url_for('admin.admin_reports'))

    html_string = render_template(
        'print_urgency.html',
        from_date=from_date,
        to_date=to_date,
        filter_department='',
        urgency_urgent=urgency_row.urgent or 0,
        urgency_planned=urgency_row.planned or 0,
        urgency_unset=urgency_row.unset or 0,
        urgency_by_dept=urgency_by_dept,
        generated_by=current_user.username,
        generated_at=datetime.now(kyiv_tz),
    )
    pdf = HTML(string=html_string).write_pdf()
    bio = BytesIO(pdf)
    bio.seek(0)
    filename = f"urgency_{from_date.strftime('%d-%m-%Y')}_{to_date.strftime('%d-%m-%Y')}.pdf"
    try:
        log_action(current_user.id, 'admin.report_urgency', 'report', None,
                   f'from={from_date} to={to_date}')
        db.session.commit()
    except Exception:
        current_app.logger.exception('Failed to log report_urgency')
    return send_file(bio, as_attachment=True, download_name=filename, mimetype='application/pdf')



# Audit Log Route
@admin_bp.route('/audit')
@role_required('admin')
def admin_audit():
    """View audit log with filters and pagination."""
    from utils import get_user_map

    # Filters
    action_filter = request.args.get('action', '').strip()
    actor_filter = request.args.get('actor', '').strip()
    from_str = request.args.get('from_date', '').strip()
    to_str = request.args.get('to_date', '').strip()
    q_search = request.args.get('q_search', '').strip()

    q = Audit.query

    if action_filter:
        q = q.filter(Audit.action.like(f'%{escape_like(action_filter)}%', escape='\\'))
    if actor_filter:
        try:
            actor_id = int(actor_filter)
            q = q.filter(Audit.actor_id == actor_id)
        except ValueError:
            pass
    if from_str:
        try:
            from_date = date.fromisoformat(from_str)
            q = q.filter(Audit.created_at >= datetime.combine(from_date, datetime.min.time()))
        except ValueError:
            pass
    if to_str:
        try:
            to_date = date.fromisoformat(to_str)
            q = q.filter(Audit.created_at < datetime.combine(to_date + timedelta(days=1), datetime.min.time()))
        except ValueError:
            pass

    if q_search:
        search_pattern = f'%{escape_like(q_search)}%'
        if q_search.isdigit():
            q = q.filter(
                (Audit.details.like(search_pattern, escape='\\')) |
                (Audit.target_type.like(search_pattern, escape='\\')) |
                (Audit.action.like(search_pattern, escape='\\')) |
                (Audit.target_id == int(q_search))
            )
        else:
            q = q.filter(
                (Audit.details.like(search_pattern, escape='\\')) |
                (Audit.target_type.like(search_pattern, escape='\\')) |
                (Audit.action.like(search_pattern, escape='\\'))
            )

    q = q.order_by(Audit.created_at.desc())

    # Pagination
    page = request.args.get('page', 1, type=int)
    per_page = clamp_per_page(request.args.get('per_page', 50), default=50)

    pagination = q.paginate(page=page, per_page=per_page, error_out=False)
    logs = pagination.items

    actions = get_distinct_audit_actions()

    user_map = get_user_map()
    users = User.query.order_by(User.username).all()

    is_htmx = request.headers.get('HX-Request') is not None

    if is_htmx:
        return render_template(
            'admin/_audit_table_partial.html',
            logs=logs,
            pagination=pagination,
            action_filter=action_filter,
            actor_filter=actor_filter,
            from_date=from_str,
            to_date=to_str,
            q_search=q_search,
            user_map=user_map,
        )

    return render_template(
        'admin_audit.html',
        logs=logs,
        pagination=pagination,
        actions=actions,
        users=users,
        user_map=user_map,
        action_filter=action_filter,
        actor_filter=actor_filter,
        from_date=from_str,
        to_date=to_str,
        q_search=q_search,
    )
